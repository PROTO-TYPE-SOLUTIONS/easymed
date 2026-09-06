from datetime import timedelta

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.template.loader import get_template
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import mixins, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.generics import ListAPIView
from rest_framework.response import Response
from rest_framework.views import APIView
from weasyprint import HTML

from company.models import Company
from customuser.models import CustomUser

from .filters import (
    IncomingItemFilter,
    InventoryFilterSearch,
    ItemFilter,
    RequisitionItemFilter,
    StockBalanceFilter,
    StockMovementFilter,
    SupplierFilter,
)
from .models import (
    Department,
    GoodsReceiptNote,
    IncomingItem,
    InsuranceItemSalePrice,
    Item,
    ItemPrice,
    ItemUnit,
    PurchaseOrder,
    PurchaseOrderItem,
    Quotation,
    QuotationItem,
    Requisition,
    RequisitionItem,
    StockBalance,
    StockLot,
    StockMovement,
    StockPolicy,
    StockReservation,
    StockTake,
    StockTakeLine,
    Supplier,
    SupplierInvoice,
    SupplierPaymentAllocation,
    SupplierPaymentReceipt,
    Unit,
)
from .serializers import (
    AllocateSupplierPaymentRequestSerializer,
    DepartmentSerializer,
    GoodsReceiptNoteSerializer,
    GoodsReceiptSerializer,
    IncomingItemSerializer,
    InsuranceItemSalePriceSerializer,
    ItemPriceSerializer,
    ItemSerializer,
    ItemUnitSerializer,
    OpeningStockSerializer,
    PurchaseOrderItemSerializer,
    PurchaseOrderSerializer,
    QuotationItemSerializer,
    QuotationSerializer,
    RequisitionItemSerializer,
    RequisitionSerializer,
    StockAdjustmentSerializer,
    StockBalanceSerializer,
    StockLotSerializer,
    StockMovementSerializer,
    StockPolicySerializer,
    StockReservationSerializer,
    StockTakeLineSerializer,
    StockTakeSerializer,
    StockTransferSerializer,
    SupplierInvoiceSerializer,
    SupplierPaymentReceiptSerializer,
    SupplierSerializer,
    UnitSerializer,
)
from .services import stock as stock_service
from .services.stock import InsufficientStock, StockError

# The dashboard still calls the stock endpoint "inventories".
InventorySerializer = StockBalanceSerializer


def _current_user(request):
    user = getattr(request, 'user', None)
    return user if user is not None and user.is_authenticated else None


class ItemUnitViewSet(viewsets.ModelViewSet):
    '''
    The pack sizes an item can be bought or sold in. One row per container
    bigger than the base unit -- a box of 12, a carton of 120.
    '''
    queryset = ItemUnit.objects.select_related('item').all()
    serializer_class = ItemUnitSerializer
    filter_backends = (DjangoFilterBackend,)
    filterset_fields = ['item', 'is_purchase_default', 'is_sale_default']


class ItemViewSet(viewsets.ModelViewSet):
    queryset = Item.objects.prefetch_related('unit_conversions').all()
    serializer_class = ItemSerializer
    filter_backends = (DjangoFilterBackend,)
    filterset_class = ItemFilter

    @action(detail=True, methods=['get'], url_path='stock')
    def stock(self, request, pk=None):
        '''Where this item is held, lot by lot.'''
        item = self.get_object()
        balances = StockBalance.objects.filter(item=item).select_related('lot', 'department', 'item')
        return Response({
            'item': item.id,
            'item_name': item.name,
            'total_on_hand': stock_service.on_hand_quantity(item),
            'available': stock_service.available_quantity(item),
            'reserved': stock_service.reserved_quantity(item),
            'balances': StockBalanceSerializer(balances, many=True).data,
        })

    @action(detail=True, methods=['get'], url_path='stock-card')
    def stock_card(self, request, pk=None):
        '''The bin card: every movement against this item, oldest first.'''
        item = self.get_object()
        department_id = request.query_params.get('department')
        department = Department.objects.filter(id=department_id).first() if department_id else None
        movements = stock_service.stock_card(item, department=department)
        return Response(StockMovementSerializer(movements, many=True).data)

    @action(detail=False, methods=['get'], url_path='export_excel')
    def export_excel(self, request):
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Items'

        columns = ['Item Code', 'Name', 'Description', 'Category', 'Unit', 'Vat Rate',
                   'Pack Name', 'Units Per Pack', 'Slow Moving Period', 'Sale Price', 'Stock Tracked']
        worksheet.append(columns)

        for item in self.get_queryset():
            # One pack size per row keeps the sheet flat. The purchase default
            # is the one buyers care about; the rest live in the API.
            pack = (
                item.unit_conversions.filter(is_purchase_default=True).first()
                or item.unit_conversions.first()
            )
            worksheet.append([
                item.item_code,
                item.name,
                item.desc,
                item.category,
                item.units_of_measure,
                str(item.vat_rate),
                pack.name if pack else '',
                pack.factor_to_base if pack else '',
                item.slow_moving_period,
                str(item.current_sale_price),
                item.is_stock_tracked,
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=items.xlsx'
        workbook.save(response)
        return response

    @action(detail=False, methods=['post'], url_path='import_excel')
    def import_excel(self, request):
        import openpyxl

        file = request.FILES.get('file')
        if not file:
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            workbook = openpyxl.load_workbook(file)
            worksheet = workbook.active

            created_count = 0
            updated_count = 0

            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue

                (item_code, name, desc, category, unit, vat_rate,
                 pack_name, units_per_pack, slow_moving_period) = row[:9]
                sale_price = row[9] if len(row) > 9 else None

                if not name or not category:
                    continue

                defaults = {
                    'item_code': item_code or '',
                    'desc': desc or '',
                    'vat_rate': vat_rate or 16.0,
                    'slow_moving_period': slow_moving_period or 90,
                }

                item, created = Item.objects.update_or_create(
                    name=name,
                    category=category,
                    units_of_measure=unit or 'units',
                    defaults=defaults,
                )

                # A pack of 1 is the base unit, which needs no conversion row.
                if pack_name and units_per_pack and int(units_per_pack) > 1:
                    ItemUnit.objects.update_or_create(
                        item=item,
                        name=str(pack_name),
                        defaults={
                            'factor_to_base': int(units_per_pack),
                            'is_purchase_default': True,
                        },
                    )

                if sale_price is not None:
                    stock_service.set_sale_price(item, sale_price, created_by=_current_user(request))

                created_count += int(created)
                updated_count += int(not created)

            return Response({
                "message": f"Successfully imported items. Created: {created_count}, Updated: {updated_count}"
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


class ItemPriceViewSet(viewsets.ModelViewSet):
    queryset = ItemPrice.objects.all().select_related('item')
    serializer_class = ItemPriceSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['item']


class UnitViewSet(viewsets.ModelViewSet):
    queryset = Unit.objects.all()
    serializer_class = UnitSerializer
    filter_backends = (DjangoFilterBackend,)
    filterset_fields = ['category']


class IncomingItemViewSet(viewsets.ModelViewSet):
    queryset = IncomingItem.objects.all().select_related('item', 'supplier', 'department')
    serializer_class = IncomingItemSerializer
    filter_backends = [InventoryFilterSearch, DjangoFilterBackend]
    filterset_class = IncomingItemFilter
    search_fields = ['lot_no', 'item__name', 'item__item_code', 'supplier__official_name']

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def destroy(self, request, *args, **kwargs):
        '''
        A posted receipt is part of the ledger's history. Deleting it would
        leave the stock it created unexplained, so reverse it instead.
        '''
        incoming_item = self.get_object()
        if incoming_item.is_posted:
            return Response(
                {"detail": "This receipt has been posted to the stock ledger. Use "
                           "/inventory/stock-movements/<id>/reverse/ or record a return to supplier."},
                status=status.HTTP_409_CONFLICT,
            )
        return super().destroy(request, *args, **kwargs)


class DepartmentViewSet(viewsets.ModelViewSet):
    queryset = Department.objects.all()
    serializer_class = DepartmentSerializer


class RequisitionViewSet(viewsets.ModelViewSet):
    queryset = Requisition.objects.all().order_by('-id')
    serializer_class = RequisitionSerializer
    filter_backends = [InventoryFilterSearch, DjangoFilterBackend]
    filterset_fields = ['requested_by', 'department']
    search_fields = [
        'requisition_number', 'requested_by__first_name', 'requested_by__last_name',
        'department__name', 'approved_by__first_name', 'approved_by__last_name',
    ]

    def _close(self, request, closed_as):
        '''
        Shared body of reject/cancel. Ending a requisition is an explicit act
        with a reason attached, so it gets its own endpoint rather than riding
        on a PATCH of a status field.
        '''
        requisition = self.get_object()
        try:
            requisition.close(
                closed_as,
                reason=request.data.get('reason', ''),
                by=_current_user(request),
            )
        except DjangoValidationError as exc:
            raise ValidationError(exc.message_dict if hasattr(exc, 'message_dict')
                                  else exc.messages)
        return Response(self.get_serializer(requisition).data)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        '''Approver declines the requisition. POST {"reason": "..."}'''
        return self._close(request, Requisition.Status.REJECTED)

    @action(detail=True, methods=['post'])
    def cancel(self, request, pk=None):
        '''Requesting side withdraws it. POST {"reason": "..."}'''
        return self._close(request, Requisition.Status.CANCELLED)

    @action(detail=True, methods=['post'])
    def reopen(self, request, pk=None):
        '''Undo a rejection or cancellation made in error.'''
        requisition = self.get_object()
        try:
            requisition.reopen()
        except DjangoValidationError as exc:
            raise ValidationError(exc.message_dict if hasattr(exc, 'message_dict')
                                  else exc.messages)
        return Response(self.get_serializer(requisition).data)


class RequisitionItemViewSet(viewsets.ModelViewSet):
    queryset = RequisitionItem.objects.all()
    serializer_class = RequisitionItemSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_class = RequisitionItemFilter

    def get_queryset(self):
        requisition_id = self.kwargs.get('requisition_pk')
        return RequisitionItem.objects.filter(requisition=requisition_id)

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['requisition_id'] = self.kwargs.get('requisition_pk')
        return context


# ---------------------------------------------------------------------------
# Stock
# ---------------------------------------------------------------------------

class StockBalanceViewSet(mixins.CreateModelMixin, viewsets.ReadOnlyModelViewSet):
    '''
    Stock on hand, one row per item / lot / location.

    Balances cannot be edited: stock is the running total of the ledger. To
    change it, post a receipt, an issue, a transfer, an adjustment or a stock
    take. POST is accepted only as manual opening stock, which is itself
    recorded as a movement.
    '''
    queryset = StockBalance.objects.select_related('item', 'lot', 'department')
    serializer_class = StockBalanceSerializer
    filter_backends = [InventoryFilterSearch, DjangoFilterBackend]
    filterset_class = StockBalanceFilter
    search_fields = ['lot__lot_number', 'item__name', 'item__item_code', 'department__name']

    def create(self, request, *args, **kwargs):
        '''
        Manual stock entry. Accepts the payload the old "Add Inventory" form
        posted, but records it as an OPENING_BALANCE movement so the quantity
        has a documented origin.
        '''
        serializer = OpeningStockSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            with transaction.atomic():
                movement = stock_service.receive(
                    item=data['item'],
                    department=data['department'],
                    quantity=data['quantity_at_hand'],
                    unit_cost=data.get('purchase_price'),
                    lot_number=data.get('lot_number', ''),
                    expiry_date=data.get('expiry_date'),
                    performed_by=_current_user(request),
                    reason=data.get('reason') or 'Opening stock entered from the dashboard',
                    source_type=StockMovement.Source.MANUAL,
                    movement_type=StockMovement.Type.OPENING_BALANCE,
                )
                if data.get('sale_price') is not None:
                    stock_service.set_sale_price(
                        data['item'], data['sale_price'], created_by=_current_user(request))
                if data.get('re_order_level') is not None:
                    StockPolicy.objects.update_or_create(
                        item=data['item'], department=data['department'],
                        defaults={'re_order_level': data['re_order_level']},
                    )
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        balance = StockBalance.objects.get(
            item=movement.item, lot=movement.lot, department=movement.department)
        return Response(StockBalanceSerializer(balance).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='summary')
    def summary(self, request):
        '''Totals per item at a location, which is the level people reorder at.'''
        queryset = self.filter_queryset(self.get_queryset())
        rows = queryset.values(
            'item', 'item__name', 'item__item_code', 'item__category', 'department', 'department__name'
        ).annotate(quantity=Sum('quantity')).order_by('item__name')
        return Response(list(rows))

    @action(detail=False, methods=['get'], url_path='valuation')
    def valuation(self, request):
        department_id = request.query_params.get('department')
        department = Department.objects.filter(id=department_id).first() if department_id else None
        result = stock_service.stock_valuation(
            department=department, category=request.query_params.get('category'))
        return Response({
            'total_value': float(result['total_value']),
            'total_units': result['total_units'],
            'lines': StockBalanceSerializer(result['lines'], many=True).data,
        })

    @action(detail=False, methods=['get'], url_path='slow-moving-items')
    def slow_moving_items(self, request):
        department_id = request.query_params.get('department')
        department = Department.objects.filter(id=department_id).first() if department_id else None
        return Response([
            {
                'item_id': row['balance'].item_id,
                'item_name': row['balance'].item.name,
                'category': row['balance'].item.category,
                'department': row['balance'].department.name,
                'quantity': row['balance'].quantity,
                'days_without_transactions': row['days_without_movement'],
                'slow_moving_period': row['balance'].item.slow_moving_period,
                'lot_number': row['balance'].lot.lot_number,
                'expiry_date': row['balance'].lot.expiry_date,
                'purchase_price': row['balance'].unit_cost,
                'sale_price': row['balance'].item.current_sale_price,
            }
            for row in stock_service.slow_moving_balances(department=department)
        ])

    @action(detail=False, methods=['get'], url_path='reorder-levels')
    def reorder_levels(self, request):
        department_id = request.query_params.get('department')
        department = Department.objects.filter(id=department_id).first() if department_id else None
        rows = stock_service.items_below_reorder_level(
            department=department, category=request.query_params.get('category'))
        return Response([
            {
                'item_id': row['item'].id,
                'item_name': row['item'].name,
                'item_code': row['item'].item_code,
                'category': row['item'].category,
                'department': row['department'].name if row['department'] else None,
                'quantity': row['quantity'],
                're_order_level': row['re_order_level'],
            }
            for row in rows
        ])


# Kept under its historical name so existing imports and routes keep working.
InventoryViewSet = StockBalanceViewSet


class StockLotViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = StockLot.objects.select_related('item', 'supplier')
    serializer_class = StockLotSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['item', 'supplier']


class StockMovementViewSet(viewsets.ReadOnlyModelViewSet):
    '''
    The ledger. Read-only over HTTP: movements are appended by the service
    layer and corrected with reversals, never edited or deleted.
    '''
    queryset = StockMovement.objects.select_related('item', 'lot', 'department', 'performed_by')
    serializer_class = StockMovementSerializer
    filter_backends = [InventoryFilterSearch, DjangoFilterBackend]
    filterset_class = StockMovementFilter
    search_fields = ['item__name', 'item__item_code', 'lot__lot_number', 'source_reference', 'reason']

    @action(detail=True, methods=['post'], url_path='reverse')
    def reverse_movement(self, request, pk=None):
        movement = self.get_object()
        reason = (request.data or {}).get('reason', '').strip()
        if not reason:
            return Response({"reason": "A reversal must carry a reason."},
                            status=status.HTTP_400_BAD_REQUEST)
        try:
            contra = stock_service.reverse(movement, reason=reason, performed_by=_current_user(request))
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(StockMovementSerializer(contra).data, status=status.HTTP_201_CREATED)


class GoodsReceiptView(APIView):
    '''
    Receive a delivery in one go: supplier invoice, goods received note and
    every line, in a single transaction.

    The browser used to fire these as three separate requests; when the lines
    failed, the invoice and GRN were already saved and stock was never posted.
    '''

    def post(self, request):
        serializer = GoodsReceiptSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        try:
            receipt = serializer.save()
        except (InsufficientStock, StockError) as exc:
            raise ValidationError({'lines': str(exc)})
        except DjangoValidationError as exc:
            raise ValidationError(exc.message_dict if hasattr(exc, 'message_dict')
                                  else exc.messages)
        return Response(
            GoodsReceiptSerializer().to_representation(receipt),
            status=status.HTTP_201_CREATED,
        )


class StockAdjustmentView(APIView):
    '''Write a reasoned correction into the ledger.'''

    def post(self, request):
        serializer = StockAdjustmentSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        lot = data.get('lot') or stock_service.get_or_create_lot(
            data['item'], data.get('lot_number', ''), data.get('expiry_date'))

        try:
            movement = stock_service.adjust(
                item=data['item'],
                lot=lot,
                department=data['department'],
                quantity_delta=data['quantity'],
                reason=data['reason'],
                movement_type=data['movement_type'],
                performed_by=_current_user(request),
            )
        except InsufficientStock as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_409_CONFLICT)
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(StockMovementSerializer(movement).data, status=status.HTTP_201_CREATED)


class StockTransferView(APIView):
    '''Move stock between locations as two balanced ledger legs.'''

    def post(self, request):
        serializer = StockTransferSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            out_movements, in_movements = stock_service.transfer(
                item=data['item'],
                from_department=data['from_department'],
                to_department=data['to_department'],
                quantity=data['quantity'],
                performed_by=_current_user(request),
                reason=data.get('reason', ''),
            )
        except InsufficientStock as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_409_CONFLICT)
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            'out': StockMovementSerializer(out_movements, many=True).data,
            'in': StockMovementSerializer(in_movements, many=True).data,
        }, status=status.HTTP_201_CREATED)


class StockPolicyViewSet(viewsets.ModelViewSet):
    queryset = StockPolicy.objects.select_related('item', 'department')
    serializer_class = StockPolicySerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['item', 'department', 'is_active']


class StockReservationViewSet(viewsets.ModelViewSet):
    queryset = StockReservation.objects.select_related('item', 'department')
    serializer_class = StockReservationSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['item', 'department', 'status']
    http_method_names = ['get', 'post', 'delete']

    def create(self, request, *args, **kwargs):
        try:
            return super().create(request, *args, **kwargs)
        except InsufficientStock as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_409_CONFLICT)
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    def destroy(self, request, *args, **kwargs):
        '''Releasing a reservation resolves it; the row stays for the audit trail.'''
        reservation = self.get_object()
        stock_service.release_reservation(reservation)
        return Response(StockReservationSerializer(reservation).data, status=status.HTTP_200_OK)


class StockTakeViewSet(viewsets.ModelViewSet):
    queryset = StockTake.objects.select_related('department').prefetch_related('lines')
    serializer_class = StockTakeSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['department', 'status']

    def perform_create(self, serializer):
        serializer.save(counted_by=_current_user(self.request))

    @action(detail=True, methods=['post'], url_path='post')
    def post_take(self, request, pk=None):
        '''Turn the count into ADJUSTMENT movements for each variance.'''
        stock_take = self.get_object()
        try:
            movements = stock_service.post_stock_take(stock_take, performed_by=_current_user(request))
        except StockError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        return Response({
            'stock_take': StockTakeSerializer(stock_take).data,
            'movements': StockMovementSerializer(movements, many=True).data,
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=['post'], url_path='populate')
    def populate(self, request, pk=None):
        '''Pre-fill the count sheet with every lot currently held at the location.'''
        stock_take = self.get_object()
        if stock_take.status != StockTake.Status.DRAFT:
            return Response({"detail": "Only a draft stock take can be populated."},
                            status=status.HTTP_400_BAD_REQUEST)

        balances = StockBalance.objects.filter(department=stock_take.department).select_related('lot')
        created = 0
        for balance in balances:
            _, was_created = StockTakeLine.objects.get_or_create(
                stock_take=stock_take, lot=balance.lot,
                defaults={'system_quantity': balance.quantity, 'counted_quantity': balance.quantity},
            )
            created += int(was_created)
        return Response({'lines_added': created,
                         'stock_take': StockTakeSerializer(stock_take).data})


class StockTakeLineViewSet(viewsets.ModelViewSet):
    queryset = StockTakeLine.objects.select_related('lot', 'lot__item', 'stock_take')
    serializer_class = StockTakeLineSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['stock_take']

    def _guard_posted(self, instance):
        if instance.stock_take.status != StockTake.Status.DRAFT:
            raise ValidationError("This stock take has already been posted.")

    def perform_update(self, serializer):
        self._guard_posted(serializer.instance)
        serializer.save()

    def perform_destroy(self, instance):
        self._guard_posted(instance)
        instance.delete()


class InventoryFilterView(ListAPIView):
    '''
    To get low quantity drugs:   GET /inventory_filter/?category=Drug&filter_type=low_quantity
    To get near-expiry drugs:    GET /inventory_filter/?category=Drug&filter_type=near_expiry
    To get expired stock:        GET /inventory_filter/?category=Drug&filter_type=expired
    '''
    serializer_class = StockBalanceSerializer

    def get_queryset(self):
        category = self.request.query_params.get('category')
        filter_type = self.request.query_params.get('filter_type')

        if not category or not filter_type:
            raise ValidationError({"error": "Both 'category' and 'filter_type' parameters are required."})

        queryset = StockBalance.objects.filter(
            item__category=category).select_related('item', 'lot', 'department')

        if filter_type == 'low_quantity':
            department_id = self.request.query_params.get('department')
            department = Department.objects.filter(id=department_id).first() if department_id else None
            rows = stock_service.items_below_reorder_level(department=department, category=category)
            pairs = [(row['item'].id, row['department'].id if row['department'] else None) for row in rows]
            if not pairs:
                return queryset.none()
            # Balances belonging to any (item, department) pair that is below level.
            from django.db.models import Q as _Q
            condition = _Q()
            for item_id, dept_id in pairs:
                condition |= _Q(item_id=item_id, department_id=dept_id)
            return queryset.filter(condition)

        if filter_type == 'near_expiry':
            today = timezone.localdate()
            horizon = today + timedelta(days=int(self.request.query_params.get('days', 90)))
            return queryset.filter(
                quantity__gt=0, lot__expiry_date__gte=today, lot__expiry_date__lte=horizon)

        if filter_type == 'expired':
            return queryset.filter(quantity__gt=0, lot__expiry_date__lt=timezone.localdate())

        raise ValidationError({
            "error": f"Invalid filter_type: {filter_type}. "
                     "Must be 'low_quantity', 'near_expiry' or 'expired'."
        })


class SupplierViewSet(viewsets.ModelViewSet):
    queryset = Supplier.objects.all()
    serializer_class = SupplierSerializer
    filter_backends = (DjangoFilterBackend,)
    filterset_class = SupplierFilter


class SupplierInvoiceViewSet(viewsets.ModelViewSet):
    serializer_class = SupplierInvoiceSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['supplier', 'purchase_order', 'status']

    def get_queryset(self):
        return SupplierInvoice.objects.all().select_related(
            'supplier', 'purchase_order', 'purchase_order__requisition'
        ).prefetch_related('incomingitem_set__goods_receipt_note')


class PurchaseOrderViewSet(viewsets.ModelViewSet):
    serializer_class = PurchaseOrderSerializer
    http_method_names = ['get', 'post', 'put', 'patch', 'delete']
    filter_backends = [InventoryFilterSearch, DjangoFilterBackend]
    search_fields = [
        'PO_number', 'ordered_by__first_name', 'ordered_by__last_name',
        'approved_by__first_name', 'approved_by__last_name',
    ]

    def get_queryset(self):
        requisition_id = self.kwargs.get('requisition_pk')
        if requisition_id:
            return PurchaseOrder.objects.filter(requisition_id=requisition_id)
        return PurchaseOrder.objects.all()

    def get_serializer_context(self):
        return {
            'request': self.request,
            'requisition_id': self.kwargs.get('requisition_pk'),
            'requested_by': self.request.user,
        }

    def perform_create(self, serializer):
        serializer.save(created_by=_current_user(self.request))

    @action(detail=False, methods=['get'])
    def all_purchase_orders(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        serializer = PurchaseOrderSerializer(queryset, many=True)
        return Response(serializer.data)


class PurchaseOrderItemViewSet(viewsets.ModelViewSet):
    serializer_class = PurchaseOrderItemSerializer
    http_method_names = ['get', 'put', 'patch']
    lookup_field = 'id'

    def get_queryset(self):
        purchase_order_id = self.kwargs.get('purchaseorder_pk')
        return PurchaseOrderItem.objects.filter(purchase_order=purchase_order_id)


class InsuranceItemSalePriceViewSet(viewsets.ModelViewSet):
    queryset = InsuranceItemSalePrice.objects.all()
    serializer_class = InsuranceItemSalePriceSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['item', 'insurance_company']


class GoodsReceiptNoteViewSet(viewsets.ModelViewSet):
    queryset = GoodsReceiptNote.objects.all()
    serializer_class = GoodsReceiptNoteSerializer


class QuotationViewSet(viewsets.ModelViewSet):
    queryset = Quotation.objects.all()
    serializer_class = QuotationSerializer


class QuotationItemViewSet(viewsets.ModelViewSet):
    queryset = QuotationItem.objects.all()
    serializer_class = QuotationItemSerializer


# ---------------------------------------------------------------------------
# PDFs
# ---------------------------------------------------------------------------

def download_requisition_pdf(request, requisition_id):
    '''
    This view gets the generated pdf and downloads it locally
    pdf accessed here http://127.0.0.1:8080/download_requisition_pdf/26/
    '''
    company = Company.objects.first()
    company_logo_url = request.build_absolute_uri(company.logo.url) if company and company.logo else None
    requisition = get_object_or_404(Requisition, pk=requisition_id)
    requisition_items = RequisitionItem.objects.filter(requisition=requisition).select_related('item')

    total_cost = 0
    for line in requisition_items:
        # Cost and quantity are both per ordering unit, so this multiplies
        # like with like whether the line is in boxes or loose units.
        total_cost += line.effective_unit_cost * (line.quantity_approved or 0)

    requester_sig_url = (
        request.build_absolute_uri(requisition.requested_by.signature.url)
        if requisition.requested_by and requisition.requested_by.signature else None
    )
    approver_sig_url = (
        request.build_absolute_uri(requisition.approved_by.signature.url)
        if requisition.approved_by and requisition.approved_by.signature else None
    )

    context = {
        'requisition': requisition,
        'requisition_items': requisition_items,
        'company': company,
        'company_logo_url': company_logo_url,
        'total_cost': total_cost,
        'requester_sig_url': requester_sig_url,
        'approver_sig_url': approver_sig_url,
    }

    html_template = get_template('requisition.html').render(context)
    pdf_file = HTML(string=html_template).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'filename="requisition_report_{requisition_id}.pdf"'
    return response


def download_purchaseorder_pdf(request, purchaseorder_id):
    '''
    Picture this, you have 1 crate of 30 eggs
    Quantity ordered is 60 eggs.
    But on the LPO pdf we want to see 2 crates. Get it?
    '''
    purchase_order = get_object_or_404(PurchaseOrder, pk=purchaseorder_id)
    purchase_order_items = PurchaseOrderItem.objects.filter(
        purchase_order=purchase_order).select_related('requisition_item', 'requisition_item__item')
    company = Company.objects.first()
    user = CustomUser.objects.first()

    item_details = []
    total_amount = 0
    for line in purchase_order_items:
        req_item = line.requisition_item
        if req_item is None:
            continue
        unit_price = req_item.effective_unit_cost
        total_price = unit_price * line.quantity_ordered
        total_amount += total_price
        item_details.append({
            'name': req_item.item.name,
            'quantity_ordered': line.quantity_ordered,
            # The supplier is quoted in the unit we ordered in -- two crates,
            # not sixty eggs -- which is what this docstring always wanted.
            'unit_label': line.unit_label,
            'base_quantity': line.base_quantity_ordered,
            'base_unit': req_item.item.units_of_measure,
            'unit_price': unit_price,
            'total_price': total_price,
        })

    company_logo_url = request.build_absolute_uri(company.logo.url) if company and company.logo else None
    creator_sig_url = (
        request.build_absolute_uri(purchase_order.created_by.signature.url)
        if purchase_order.created_by and purchase_order.created_by.signature else None
    )
    approver_sig_url = (
        request.build_absolute_uri(purchase_order.approved_by.signature.url)
        if purchase_order.approved_by and purchase_order.approved_by.signature else None
    )

    context = {
        'purchaseorder': purchase_order,
        'item_details': item_details,
        'total_amount': total_amount,
        'company': company,
        'company_logo_url': company_logo_url,
        'user': user,
        'creator_sig_url': creator_sig_url,
        'approver_sig_url': approver_sig_url,
    }

    html_template = get_template('purchase_order_note.html').render(context)
    pdf_file = HTML(string=html_template).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'filename="purchase_order_report_{purchaseorder_id}.pdf"'
    return response


def download_goods_receipt_note_pdf(request, purchase_order_id):
    incoming_items = IncomingItem.objects.filter(
        purchase_order_id=purchase_order_id).select_related('item', 'supplier', 'goods_receipt_note')
    company = Company.objects.first()

    goods_receipt_note = incoming_items.first().goods_receipt_note if incoming_items.exists() else None
    grn_number = goods_receipt_note.grn_number if goods_receipt_note else "N/A"

    item_details = []
    total_price_before_vat = 0
    total_vat = 0
    total_amount_after_vat = 0

    for line in incoming_items:
        amount_before_vat = line.line_total
        vat_amount = amount_before_vat * (line.item.vat_rate / 100)
        amount_with_vat = amount_before_vat + vat_amount

        total_price_before_vat += amount_before_vat
        total_vat += vat_amount
        total_amount_after_vat += amount_with_vat
        item_details.append({
            'supplier': line.supplier,
            'item_code': line.item.item_code,
            'lot_number': line.lot_no,
            'item_name': line.item.name,
            'quantity_received': line.quantity,
            'quantity_unit': line.item_unit.name if line.item_unit_id else (line.item.units_of_measure or 'units'),
            'base_units': line.base_units,
            'unit_price': line.purchase_price,
            'amount_before_vat': amount_before_vat,
            'vat_amount': vat_amount,
            'amount_with_vat': amount_with_vat,
            'expiry_date': line.expiry_date,
        })

    company_logo_url = request.build_absolute_uri(company.logo.url) if company and company.logo else None

    context = {
        'incoming_items': incoming_items,
        'company': company,
        'company_logo_url': company_logo_url,
        'grn_number': grn_number,
        'item_details': item_details,
        'total_price_before_vat': total_price_before_vat,
        'total_vat': total_vat,
        'total_amount_after_vat': total_amount_after_vat,
    }

    html_template = get_template('goods_receipt_note.html').render(context)
    pdf_file = HTML(string=html_template).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="incoming_items.pdf"'
    return response


def download_supplier_invoice_pdf(request, supplier_id):
    supplier = get_object_or_404(Supplier, pk=supplier_id)
    supplier_invoices = SupplierInvoice.objects.filter(supplier=supplier).prefetch_related('incomingitem_set')
    incoming_items = IncomingItem.objects.filter(supplier_invoice__supplier=supplier)
    company = Company.objects.first()

    company_logo_url = request.build_absolute_uri(company.logo.url) if company and company.logo else None

    context = {
        'supplier': supplier,
        'supplier_invoices': supplier_invoices,
        'company': company,
        'company_logo_url': company_logo_url,
        'incoming_items': incoming_items,
    }

    html_template = get_template('supplier_invoice.html').render(context)
    pdf_file = HTML(string=html_template).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'filename="supplier_invoice_report_{supplier_id}.pdf"'
    return response


class AllocateSupplierPaymentView(APIView):
    """
    Allocate a payment to supplier invoices, creating a SupplierPaymentReceipt and allocations.
    """

    def post(self, request, *args, **kwargs):
        from billing.models import SubAccount

        req_ser = AllocateSupplierPaymentRequestSerializer(data=request.data)
        req_ser.is_valid(raise_exception=True)
        data = req_ser.validated_data

        supplier_id = data['supplier_id']
        invoice_ids = data['invoice_ids']
        amount = data['amount']

        invoices = SupplierInvoice.objects.filter(id__in=invoice_ids, supplier_id=supplier_id)
        if not invoices.exists():
            return Response({"detail": "No invoices found for the selected supplier."},
                            status=status.HTTP_400_BAD_REQUEST)

        try:
            sub_account = SubAccount.objects.select_related('payment_mode').get(id=data['sub_account'])
        except SubAccount.DoesNotExist:
            return Response({"detail": "Invalid sub account."}, status=status.HTTP_400_BAD_REQUEST)

        account_balance = sub_account.balance
        if account_balance <= 0:
            return Response(
                {"detail": f"Sub account '{sub_account.name}' has zero balance. Cannot make payment."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if amount > account_balance:
            return Response(
                {"detail": f"Insufficient funds. Sub account '{sub_account.name}' has a balance of "
                           f"{account_balance:.2f} but the payment amount is {amount:.2f}."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        with transaction.atomic():
            receipt = SupplierPaymentReceipt.objects.create(
                supplier_id=supplier_id,
                sub_account=sub_account,
                payment_mode=sub_account.payment_mode,
                total_amount=amount,
                reference_number=data['reference_number'],
                payment_date=data.get('payment_date'),
            )

            remaining = float(amount)

            for invoice in invoices.order_by('date_created'):
                if remaining <= 0:
                    break

                already_paid = float(
                    invoice.payment_allocations.aggregate(total=Sum('amount_applied'))['total'] or 0)
                outstanding = float(invoice.amount) - already_paid
                if outstanding <= 0:
                    continue

                apply_now = min(remaining, outstanding)
                if apply_now > 0:
                    SupplierPaymentAllocation.objects.create(
                        receipt=receipt,
                        supplier_invoice=invoice,
                        amount_applied=apply_now,
                    )
                    remaining -= apply_now

                    new_outstanding = outstanding - apply_now
                    invoice.status = 'paid' if new_outstanding <= 0.01 else 'partial'
                    invoice.save(update_fields=['status'])

            return Response(SupplierPaymentReceiptSerializer(receipt).data, status=status.HTTP_201_CREATED)


class SupplierPaymentReceiptViewSet(viewsets.ReadOnlyModelViewSet):
    """
    ViewSet for listing and retrieving supplier payment receipts.
    """
    queryset = SupplierPaymentReceipt.objects.all().select_related(
        'supplier', 'payment_mode', 'sub_account', 'sub_account__main_account'
    ).order_by('-created_at')
    serializer_class = SupplierPaymentReceiptSerializer
    filter_backends = [DjangoFilterBackend]
    filterset_fields = ['supplier', 'payment_mode', 'payment_date']


def download_supplier_payment_receipt_pdf(request, receipt_id):
    receipt = get_object_or_404(
        SupplierPaymentReceipt.objects.select_related(
            'supplier', 'sub_account', 'sub_account__main_account', 'payment_mode'
        ),
        pk=receipt_id,
    )
    company = Company.objects.first()
    company_logo_url = (
        request.build_absolute_uri(company.logo.url)
        if company and getattr(company, 'logo', None) and company.logo
        else None
    )

    allocations = receipt.allocations.select_related(
        'supplier_invoice', 'supplier_invoice__supplier').all()
    total_invoiced = sum(a.supplier_invoice.amount for a in allocations)

    html_template = get_template('supplier_payment_receipt.html').render({
        'company': company,
        'company_logo_url': company_logo_url,
        'receipt': receipt,
        'allocations': allocations,
        'total_invoiced': total_invoiced,
    })

    pdf_file = HTML(string=html_template).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'filename="supplier_payment_receipt_{receipt.id}.pdf"'
    return response
